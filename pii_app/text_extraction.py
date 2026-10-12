"""Text extraction for PDF, DOCX, PPTX, TXT, PNG, and JPG artifacts.

Each extractor takes raw file bytes and returns plain text. PDFs fall back
to OCR per-page when a page yields no selectable text (i.e. it is likely a
scanned image), so scanned PDFs are supported on a best-effort basis.
"""

import io
from dataclasses import dataclass

import pymupdf
import pytesseract
from docx import Document
from PIL import Image
from pptx import Presentation


@dataclass
class ExtractionResult:
    text: str
    used_ocr: bool = False
    pages_or_slides: int = 0
    warnings: list = None

    def __post_init__(self):
        if self.warnings is None:
            self.warnings = []


def extract_from_txt(file_bytes: bytes) -> ExtractionResult:
    text = file_bytes.decode("utf-8", errors="replace")
    return ExtractionResult(text=text)


def extract_from_docx(file_bytes: bytes) -> ExtractionResult:
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return ExtractionResult(text="\n".join(parts), pages_or_slides=1)


def extract_from_pptx(file_bytes: bytes) -> ExtractionResult:
    prs = Presentation(io.BytesIO(file_bytes))
    parts = []
    slide_count = 0

    for slide in prs.slides:
        slide_count += 1
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line:
                        parts.append(line)

    return ExtractionResult(text="\n".join(parts), pages_or_slides=slide_count)


def extract_from_image(file_bytes: bytes) -> ExtractionResult:
    image = Image.open(io.BytesIO(file_bytes))
    text = pytesseract.image_to_string(image)
    return ExtractionResult(text=text, used_ocr=True, pages_or_slides=1)


def extract_from_pdf(file_bytes: bytes, ocr_min_chars: int = 20) -> ExtractionResult:
    """Extract text from a PDF, OCR-ing any page with little/no selectable text."""
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    parts = []
    used_ocr = False
    warnings = []

    for page_index, page in enumerate(doc):
        page_text = page.get_text().strip()

        if len(page_text) < ocr_min_chars:
            # Likely a scanned page: render to an image and OCR it.
            try:
                pixmap = page.get_pixmap(dpi=300)
                img_bytes = pixmap.tobytes("png")
                image = Image.open(io.BytesIO(img_bytes))
                ocr_text = pytesseract.image_to_string(image)
                if len(ocr_text.strip()) > len(page_text):
                    page_text = ocr_text
                    used_ocr = True
            except Exception as exc:  # pragma: no cover - defensive, OCR env issues
                warnings.append(f"OCR failed on page {page_index + 1}: {exc}")

        parts.append(page_text)

    doc.close()
    return ExtractionResult(
        text="\n".join(parts),
        used_ocr=used_ocr,
        pages_or_slides=len(parts),
        warnings=warnings,
    )


EXTENSION_TO_EXTRACTOR = {
    "txt": extract_from_txt,
    "docx": extract_from_docx,
    "pptx": extract_from_pptx,
    "pdf": extract_from_pdf,
    "png": extract_from_image,
    "jpg": extract_from_image,
    "jpeg": extract_from_image,
}


def extract_text(filename: str, file_bytes: bytes) -> ExtractionResult:
    extension = filename.rsplit(".", 1)[-1].lower()
    extractor = EXTENSION_TO_EXTRACTOR.get(extension)
    if extractor is None:
        raise ValueError(f"No extractor registered for extension '.{extension}'")
    return extractor(file_bytes)
