"""Generate synthetic sample artifacts (DOCX, PPTX, PNG) for the demo.

Run once from the project root: python sample_data/generate_samples.py
All data here is fictional, created for this case study only.
"""

import os

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))


def generate_docx():
    doc = Document()
    doc.add_heading("Employee Onboarding Form", level=1)

    doc.add_paragraph("Full Name: Neha Sharma")
    doc.add_paragraph("Employee ID: EMP-30077")
    doc.add_paragraph("Email: neha.sharma@examplecorp.com")
    doc.add_paragraph("Phone: +91 9811223344")
    doc.add_paragraph("Date of Birth: 05/11/1995")
    doc.add_paragraph("PAN: LMNOP9876Q")
    doc.add_paragraph("Aadhaar: 1234 5678 9012")
    doc.add_paragraph(
        "Address: 7 Rosewood Enclave, Andheri West, Mumbai, Maharashtra, 400058"
    )
    doc.add_paragraph(
        "Reporting manager Arjun Verma confirmed Neha Sharma's start date "
        "for the Bengaluru office transfer."
    )

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Emergency Contact"
    table.rows[0].cells[1].text = "Arjun Verma"
    table.rows[1].cells[0].text = "Emergency Phone"
    table.rows[1].cells[1].text = "9900112233"

    out_path = os.path.join(HERE, "sample_onboarding_form.docx")
    doc.save(out_path)
    print(f"Wrote {out_path}")


def generate_pptx():
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]

    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Vendor Risk Review"
    body = slide.placeholders[1].text_frame
    body.text = "Point of contact: Vikram Rao (vikram.rao@examplecorp.com)"
    p = body.add_paragraph()
    p.text = "Employee ID: EMP-44210, Phone: +91 9654321098"

    slide2 = prs.slides.add_slide(slide_layout)
    slide2.shapes.title.text = "Audit Notes"
    body2 = slide2.placeholders[1].text_frame
    body2.text = "PAN on file: STUVX1122Y"
    p2 = body2.add_paragraph()
    p2.text = "Aadhaar on file: 9012 3456 7890"
    p3 = body2.add_paragraph()
    p3.text = "DOB: 19/02/1990, reviewed by Kavya Reddy"

    out_path = os.path.join(HERE, "sample_vendor_review.pptx")
    prs.save(out_path)
    print(f"Wrote {out_path}")


def generate_image():
    """Render a simple text-on-white image to exercise OCR extraction."""
    width, height = 900, 400
    image = Image.new("RGB", (width, height), color="white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 28)
    except OSError:
        font = ImageFont.load_default()

    lines = [
        "SCANNED ID CARD (SAMPLE)",
        "Name: Tanvi Joshi",
        "Employee ID: EMP-50521",
        "Email: tanvi.joshi@examplecorp.com",
        "Phone: +91 9090909090",
        "DOB: 30/06/1993",
    ]

    y = 30
    for line in lines:
        draw.text((40, y), line, fill="black", font=font)
        y += 55

    out_path = os.path.join(HERE, "sample_scanned_id.png")
    image.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    generate_docx()
    generate_pptx()
    generate_image()
